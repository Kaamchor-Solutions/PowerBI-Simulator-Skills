"""Smoke and contract tests for the constrained helper (not Power BI parity tests)."""
import csv
from contextlib import redirect_stderr
import io
import json
import os
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from decimal import (
    Decimal,
    Inexact,
    ROUND_05UP,
    ROUND_CEILING,
    ROUND_DOWN,
    ROUND_FLOOR,
    ROUND_HALF_DOWN,
    ROUND_HALF_EVEN,
    ROUND_HALF_UP,
    ROUND_UP,
    Rounded,
    Subnormal,
    localcontext,
)
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import analyze  # noqa: E402
import bi_report  # noqa: E402

DEMO = ROOT / "examples" / "demo_sales"
TABLES = {
    "Orders": {"file": str(DEMO / "orders.csv")},
    "Customers": {"file": str(DEMO / "customers.csv")},
    "Products": {"file": str(DEMO / "products.csv")},
}
RELATIONSHIPS = [
    {"one_table": "Customers", "one_column": "CustomerID", "many_table": "Orders", "many_column": "CustomerID"},
    {"one_table": "Products", "one_column": "ProductID", "many_table": "Orders", "many_column": "ProductID"},
]


class AnalyzeTests(unittest.TestCase):
    def test_inventory_reports_shape_and_missingness(self):
        result = analyze.run_request({"action": "inventory", "tables": TABLES})
        self.assertTrue(result["valid"])
        self.assertEqual(result["tables"]["Orders"]["rows"], 10)
        fields = {field["name"]: field for field in result["tables"]["Orders"]["fields"]}
        self.assertEqual(fields["Quantity"]["inferred_type"], "integer")
        self.assertEqual(fields["DiscountPct"]["missing"], 0)

    def test_declared_star_relationships_pass(self):
        result = analyze.run_request({"action": "model_audit", "tables": TABLES, "relationships": RELATIONSHIPS})
        self.assertTrue(result["valid"])
        self.assertEqual(len(result["relationships"]), 2)
        self.assertEqual(sum(item["orphan_count"] for item in result["relationships"]), 0)

    def test_orphans_are_reported_and_model_is_invalid(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "orders.csv"
            with path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(["OrderID", "CustomerID"])
                writer.writerow(["O1", "missing-customer"])
            tables = {"Customers": TABLES["Customers"], "Orders": {"file": str(path)}}
            result = analyze.run_request({
                "action": "model_audit",
                "tables": tables,
                "relationships": [{"one_table": "Customers", "one_column": "CustomerID", "many_table": "Orders", "many_column": "CustomerID"}],
            })
        self.assertFalse(result["valid"])
        self.assertEqual(result["relationships"][0]["orphan_count"], 1)

    def test_model_audit_samples_are_opt_in_and_limited(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            one = folder / "one.csv"
            many = folder / "many.csv"
            one.write_text("Key,Label\nDUPLICATE-SECRET,A\nDUPLICATE-SECRET,B\n,Null-Key\n", encoding="utf-8")
            many.write_text("ForeignKey,Label\nORPHAN-SECRET-A,A\nORPHAN-SECRET-B,B\n,Null-Key\n", encoding="utf-8")
            request = {
                "action": "model_audit",
                "tables": {"One": {"file": str(one)}, "Many": {"file": str(many)}},
                "relationships": [{
                    "one_table": "One", "one_column": "Key",
                    "many_table": "Many", "many_column": "ForeignKey",
                }],
            }
            hidden = analyze.run_request(request)
            visible = analyze.run_request({**request, "include_samples": True, "sample_limit": 1})

        hidden_relationship = hidden["relationships"][0]
        self.assertEqual(hidden_relationship["one_key_nulls"], 1)
        self.assertEqual(hidden_relationship["duplicate_one_key_count"], 1)
        self.assertEqual(hidden_relationship["orphan_count"], 2)
        self.assertEqual(hidden_relationship["many_key_nulls"], 1)
        self.assertEqual(hidden_relationship["duplicate_one_key_samples"], [])
        self.assertEqual(hidden_relationship["orphan_samples"], [])
        self.assertNotIn("SECRET", analyze._json_dumps(hidden))
        hidden_markdown = analyze.format_markdown(hidden)
        self.assertIn("## Relationship diagnostics", hidden_markdown)
        self.assertIn("Null one-side keys", hidden_markdown)
        self.assertIn("Duplicate one-side keys", hidden_markdown)
        self.assertIn("Null many-side keys", hidden_markdown)
        self.assertIn("Orphan many-side keys", hidden_markdown)
        self.assertIn("| False | 1 | 1 | — | 1 | 2 | — |", hidden_markdown)
        self.assertNotIn("SECRET", hidden_markdown)
        visible_relationship = visible["relationships"][0]
        self.assertEqual(visible_relationship["duplicate_one_key_samples"], ["DUPLICATE-SECRET"])
        self.assertEqual(visible_relationship["orphan_samples"], ["ORPHAN-SECRET-A"])
        self.assertNotIn("ORPHAN-SECRET-B", analyze._json_dumps(visible))
        visible_markdown = analyze.format_markdown(visible)
        self.assertIn("DUPLICATE-SECRET", visible_markdown)
        self.assertIn("ORPHAN-SECRET-A", visible_markdown)
        self.assertNotIn("ORPHAN-SECRET-B", visible_markdown)

    def test_inventory_samples_are_opt_in_and_limited(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "source.csv"
            path.write_text("Name\nFIRST-SECRET\nSECOND-SECRET\n", encoding="utf-8")
            request = {"action": "inventory", "tables": {"Sales": {"file": str(path)}}}
            hidden = analyze.run_request(request)
            visible = analyze.run_request({**request, "include_samples": True, "sample_limit": 1})

        hidden_result = analyze._json_dumps(hidden)
        self.assertEqual(hidden["tables"]["Sales"]["fields"][0]["sample_values"], [])
        self.assertNotIn("SECRET", hidden_result)
        self.assertNotIn("SECRET", analyze.format_markdown(hidden))
        visible_result = analyze._json_dumps(visible)
        self.assertEqual(visible["tables"]["Sales"]["fields"][0]["sample_values"], ["FIRST-SECRET"])
        self.assertIn("FIRST-SECRET", visible_result)
        self.assertNotIn("SECOND-SECRET", visible_result)

    def test_invalid_sample_policy_is_rejected_before_file_io(self):
        request = {
            "action": "inventory",
            "tables": {"Sales": {"file": "/path/that/must/not/be/opened.csv"}},
            "include_samples": "yes",
        }
        with self.assertRaisesRegex(analyze.RequestError, "include_samples.*boolean"):
            analyze.run_request(request)
        request["include_samples"] = True
        request["sample_limit"] = True
        with self.assertRaisesRegex(analyze.RequestError, "sample_limit.*0.*5"):
            analyze.run_request(request)
        request["sample_limit"] = 6
        with self.assertRaisesRegex(analyze.RequestError, "sample_limit.*0.*5"):
            analyze.run_request(request)

    def test_grouped_sales_aggregate_matches_fixture(self):
        request = {
            "action": "aggregate",
            "tables": TABLES,
            "relationships": RELATIONSHIPS,
            "fact_table": "Orders",
            "group_by": [{"table": "Products", "column": "Category"}],
            "measures": [
                {"name": "Net Sales", "op": "sum_expression", "expression": "Quantity * UnitPrice * (1 - DiscountPct)"},
                {"name": "Order Count", "op": "distinct_count", "table": "Orders", "column": "OrderID"},
            ],
        }
        with patch.object(analyze, "_compile_expression", wraps=analyze._compile_expression) as compile_expression:
            result = analyze.run_request(request)
        self.assertEqual(compile_expression.call_count, 1)
        self.assertTrue(result["valid"])
        self.assertEqual(result["totals"]["Net Sales"], 1705)
        self.assertEqual(result["totals"]["Order Count"], 10)
        by_category = {row["group"]["Products.Category"]: row["measures"]["Net Sales"] for row in result["groups"]}
        self.assertEqual(by_category, {"Accessories": 635, "Computers": 595, "Software": 475})

    def test_all_aggregate_operations_and_in_neq_filters(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "measures.csv"
            path.write_text("Amount,Group\n1,A\n2,B\n2,B\n,B\n", encoding="utf-8")
            measures = [
                {"name": "Rows", "op": "count_rows"},
                {"name": "Known", "op": "count_nonblank", "table": "Sales", "column": "Amount"},
                {"name": "Distinct", "op": "distinct_count", "table": "Sales", "column": "Amount"},
                {"name": "Sum", "op": "sum", "table": "Sales", "column": "Amount"},
                {"name": "Mean", "op": "mean", "table": "Sales", "column": "Amount"},
                {"name": "Min", "op": "min", "table": "Sales", "column": "Amount"},
                {"name": "Max", "op": "max", "table": "Sales", "column": "Amount"},
                {"name": "Expression", "op": "sum_expression", "expression": "Amount * 2"},
            ]
            common = {
                "action": "aggregate",
                "tables": {"Sales": {"file": str(path)}},
                "fact_table": "Sales",
                "measures": measures,
            }
            for filter_item in (
                {"table": "Sales", "column": "Group", "op": "in", "value": ["B"]},
                {"table": "Sales", "column": "Group", "op": "neq", "value": "A"},
            ):
                result = analyze.run_request({**common, "filters": [filter_item]})
                self.assertTrue(result["valid"])
                self.assertEqual(result["rows_used"], 3)
                self.assertEqual(result["totals"], {
                    "Rows": 3,
                    "Known": 2,
                    "Distinct": 1,
                    "Sum": 4,
                    "Mean": 2,
                    "Min": 2,
                    "Max": 2,
                    "Expression": 8,
                })
                self.assertEqual(result["measure_audit"]["Expression"]["blank_or_skipped"], 1)

    def test_tsv_and_tab_sources_are_read(self):
        with tempfile.TemporaryDirectory() as temp:
            results = {}
            for suffix in (".tsv", ".tab"):
                path = Path(temp) / f"source{suffix}"
                path.write_text("Key\tAmount\nA\t1.25\n", encoding="utf-8")
                results[suffix] = analyze.run_request({
                    "action": "inventory",
                    "tables": {"Sales": {"file": str(path)}},
                    "include_samples": True,
                })
        for result in results.values():
            self.assertTrue(result["valid"])
            fields = result["tables"]["Sales"]["fields"]
            self.assertEqual(fields[1]["inferred_type"], "decimal")
            self.assertEqual(fields[1]["sample_values"], ["1.25"])

    def test_decimal_aggregate_is_preserved_as_an_exact_json_number(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "amounts.csv"
            path.write_text("Amount\n123456789012345.6789012345\n0.0000000001\n", encoding="utf-8")
            result = analyze.run_request({
                "action": "aggregate",
                "tables": {"Sales": {"file": str(path)}},
                "fact_table": "Sales",
                "measures": [{"name": "Precise Total", "op": "sum", "table": "Sales", "column": "Amount"}],
            })
        expected = Decimal("123456789012345.6789012346")
        self.assertEqual(result["totals"]["Precise Total"], expected)
        serialized = analyze._json_dumps(result, indent=2)
        self.assertIn('"Precise Total": 123456789012345.6789012346', serialized)
        self.assertNotIn('"Precise Total": "123456789012345.6789012346"', serialized)
        decoded = json.loads(serialized, parse_float=Decimal)
        self.assertEqual(decoded["totals"]["Precise Total"], expected)

    def test_sum_is_exact_beyond_the_default_decimal_context(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "amounts.csv"
            path.write_text(
                "Amount\n1234567890123456789012345678901234567890\n"
                "0.00000000000000000000000000001\n",
                encoding="utf-8",
            )
            result = analyze.run_request({
                "action": "aggregate",
                "tables": {"Sales": {"file": str(path)}},
                "fact_table": "Sales",
                "measures": [{"name": "Total", "op": "sum", "table": "Sales", "column": "Amount"}],
            })
        self.assertEqual(
            result["totals"]["Total"],
            Decimal("1234567890123456789012345678901234567890.00000000000000000000000000001"),
        )

    def test_expression_multiplication_and_decimal_literal_remain_exact(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "amounts.csv"
            path.write_text("Amount\n2\n", encoding="utf-8")
            result = analyze.run_request({
                "action": "aggregate",
                "tables": {"Sales": {"file": str(path)}},
                "fact_table": "Sales",
                "measures": [
                    {"name": "Large product", "op": "sum_expression", "expression": "Amount * 12345678901234567890123456789"},
                    {"name": "Long literal", "op": "sum_expression", "expression": "Amount * 0.12345678901234567890123456789"},
                ],
            })
        self.assertEqual(result["totals"]["Large product"], Decimal("24691357802469135780246913578"))
        self.assertEqual(result["totals"]["Long literal"], Decimal("0.24691357802469135780246913578"))

    def test_cli_preserves_decimal_json_filter_tokens(self):
        token = "0.12345678901234567890123456789"
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            data = folder / "amounts.csv"
            data.write_text(f"Amount\n{token}\n", encoding="utf-8")
            config = folder / "request.json"
            config.write_text(
                '{"action":"aggregate","tables":{"Sales":{"file":'
                + json.dumps(str(data))
                + '}},"fact_table":"Sales","filters":[{"table":"Sales",'
                '"column":"Amount","op":"eq","value":'
                + token
                + '}],"measures":[{"name":"Total","op":"sum",'
                '"table":"Sales","column":"Amount"}]}',
                encoding="utf-8",
            )
            completed = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "analyze.py"), "--config", str(config)],
                capture_output=True,
                text=True,
                check=False,
            )
        self.assertEqual(completed.returncode, 0, completed.stderr)
        result = json.loads(completed.stdout, parse_float=Decimal)
        self.assertEqual(result["rows_used"], 1)
        self.assertEqual(result["totals"]["Total"], Decimal(token))

    def test_report_builder_preserves_decimal_json_filter_tokens(self):
        token = "0.12345678901234567890123456789"
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            data = folder / "amounts.csv"
            data.write_text(f"Amount\n{token}\n", encoding="utf-8")
            config = folder / "report.json"
            config.write_text(
                '{"analyses":[{"name":"Exact","request":{"action":"aggregate",'
                '"tables":{"Sales":{"file":' + json.dumps(str(data)) + '}}'
                ',"fact_table":"Sales","filters":[{"table":"Sales",'
                '"column":"Amount","op":"eq","value":' + token + '}],'
                '"measures":[{"name":"Total","op":"sum","table":"Sales",'
                '"column":"Amount"}]}}]}',
                encoding="utf-8",
            )
            output = folder / "report.md"
            completed = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "bi_report.py"), "--config", str(config), "--output", str(output)],
                capture_output=True,
                text=True,
                check=False,
            )
            report = output.read_text(encoding="utf-8") if output.exists() else ""
        self.assertEqual(completed.returncode, 0, completed.stderr)
        self.assertIn(token, report)

    def test_nonstandard_json_constant_is_rejected(self):
        with self.assertRaises(analyze.RequestError):
            analyze.load_json(io.StringIO('{"value": NaN}'))

    def test_duplicate_json_object_members_are_rejected_at_every_depth(self):
        fixtures = (
            ('{"action":"inventory","action":"aggregate"}', "action"),
            ('{"tables":{"Sales":{"file":"one.csv","file":"two.csv"}}}', "file"),
        )
        for text, duplicate in fixtures:
            with self.subTest(duplicate=duplicate):
                with self.assertRaisesRegex(analyze.RequestError, f"Duplicate JSON object key {duplicate!r}"):
                    analyze.load_json(io.StringIO(text))

    def test_duplicate_json_object_member_is_a_clean_cli_error(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Path(temp) / "request.json"
            config.write_text('{"action":"inventory","action":"aggregate","tables":{}}', encoding="utf-8")
            completed = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "analyze.py"), "--config", str(config)],
                capture_output=True, text=True, check=False,
            )
        self.assertEqual(completed.returncode, 2)
        self.assertIn("Duplicate JSON object key 'action'", completed.stderr)
        self.assertNotIn("Traceback", completed.stderr)
        payload = json.loads(completed.stdout)
        self.assertFalse(payload["valid"])
        self.assertEqual(payload["error_code"], "duplicate_json_key")

    def test_duplicate_json_object_member_is_a_clean_report_cli_error(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Path(temp) / "report.json"
            output = Path(temp) / "report.md"
            config.write_text('{"analyses":[],"analyses":[]}', encoding="utf-8")
            completed = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "bi_report.py"), "--config", str(config), "--output", str(output)],
                capture_output=True, text=True, check=False,
            )
        self.assertEqual(completed.returncode, 2)
        self.assertIn("Duplicate JSON object key 'analyses'", completed.stderr)
        self.assertNotIn("Traceback", completed.stderr)
        self.assertFalse(output.exists())

    def test_malformed_json_cli_includes_stable_invalid_json_code(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Path(temp) / "request.json"
            config.write_text('{"action":', encoding="utf-8")
            completed = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "analyze.py"), "--config", str(config)],
                capture_output=True, text=True, check=False,
            )
        self.assertEqual(completed.returncode, 2)
        self.assertIn("Error:", completed.stderr)
        self.assertNotIn("Traceback", completed.stderr)
        self.assertEqual(json.loads(completed.stdout)["error_code"], "invalid_json")

    def test_extreme_json_exponent_becomes_a_clear_cli_error(self):
        with tempfile.TemporaryDirectory() as temp:
            config = Path(temp) / "request.json"
            config.write_text(
                '{"action":"inventory","tables":{"Sales":{"file":"unused.csv"}},'
                '"unused":1e999999999999999999999999999999999999999999999}',
                encoding="utf-8",
            )
            completed = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "analyze.py"), "--config", str(config)],
                capture_output=True,
                text=True,
                check=False,
            )
        self.assertEqual(completed.returncode, 2)
        self.assertIn("outside the supported Decimal range", completed.stderr)
        self.assertNotIn("Traceback", completed.stderr)

    def test_long_json_integer_is_preserved_without_python_conversion_limit(self):
        digits = "9" * 5_000
        parsed = analyze.load_json(io.StringIO('{"value":' + digits + '}'))
        self.assertEqual(parsed["value"], Decimal(digits))

    def test_exact_arithmetic_precision_limit_is_enforced(self):
        with self.assertRaisesRegex(analyze.RequestError, "100,000 digits"):
            analyze._decimal_add(Decimal("1e100001"), Decimal(1))

    def test_decimal_arithmetic_isolated_from_ambient_precision_and_traps(self):
        large = Decimal("12345678901234567890123456789")
        tiny = Decimal("1e-1000000000000000000")
        with localcontext() as ambient:
            ambient.prec = 2
            ambient.traps[Inexact] = True
            ambient.traps[Rounded] = True
            ambient.traps[Subnormal] = True
            self.assertEqual(analyze._decimal_add(large, Decimal(1)), Decimal("12345678901234567890123456790"))
            self.assertEqual(analyze._decimal_subtract(large, Decimal(1)), Decimal("12345678901234567890123456788"))
            self.assertEqual(analyze._decimal_multiply(large, Decimal(2)), Decimal("24691357802469135780246913578"))
            self.assertEqual(analyze._decimal_sum([large, Decimal(1), Decimal(1)]), Decimal("12345678901234567890123456791"))
            self.assertEqual(analyze._decimal_add(tiny, tiny), Decimal("2e-1000000000000000000"))
            self.assertEqual(
                analyze._decimal_divide(Decimal(1), Decimal(3)),
                Decimal("0.33333333333333333333333333333333333333333333333333"),
            )

    def test_division_is_independent_of_all_ambient_rounding_modes(self):
        expected = Decimal("0.33333333333333333333333333333333333333333333333333")
        rounding_modes = (
            ROUND_05UP, ROUND_CEILING, ROUND_DOWN, ROUND_FLOOR,
            ROUND_HALF_DOWN, ROUND_HALF_EVEN, ROUND_HALF_UP, ROUND_UP,
        )
        for rounding in rounding_modes:
            with self.subTest(rounding=rounding):
                with localcontext() as ambient:
                    ambient.prec = 3
                    ambient.rounding = rounding
                    actual = analyze._decimal_divide(Decimal(1), Decimal(3))
                self.assertEqual(actual, expected)

    def test_report_builder_mixes_success_and_request_error_with_limitations(self):
        request = {
            "action": "aggregate",
            "tables": TABLES,
            "relationships": RELATIONSHIPS,
            "fact_table": "Orders",
            "measures": [{"name": "Net Sales", "op": "sum_expression", "expression": "Quantity * UnitPrice * (1 - DiscountPct)"}],
        }
        report = bi_report.build_report({
            "title": "Contract test",
            "analyses": [
                {"name": "Valid analysis", "request": request},
                {"name": "Invalid analysis", "request": {**request, "action": "unsupported"}},
            ],
        })
        self.assertIn("## Valid analysis", report)
        self.assertIn("| Net Sales | 1705 |", report)
        self.assertIn("## Invalid analysis", report)
        self.assertIn("**Did not run:**", report)
        self.assertIn("## Limitations", report)

    def test_report_cli_strict_mode_returns_nonzero_and_writes_mixed_report(self):
        request = {
            "analyses": [
                {"name": "Valid analysis", "request": {"action": "inventory", "tables": TABLES}},
                {"name": "Invalid analysis", "request": {"action": "unsupported"}},
            ]
        }
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            config = folder / "report.json"
            config.write_text(json.dumps(request), encoding="utf-8")
            output = folder / "report.md"
            completed = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "bi_report.py"), "--config", str(config),
                 "--output", str(output), "--fail-on-invalid"],
                capture_output=True,
                text=True,
                check=False,
            )
            report = output.read_text(encoding="utf-8") if output.exists() else ""
        self.assertEqual(completed.returncode, 2)
        self.assertIn("REPORT STATUS: PARTIAL", report)
        self.assertIn("## Valid analysis", report)
        self.assertIn("## Invalid analysis", report)
        self.assertNotIn("Traceback", completed.stderr)

    def test_report_builder_reports_invalid_error_skipped_and_empty_statuses(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            one_side = folder / "one.csv"
            one_side.write_text("Key\nA\nA\n", encoding="utf-8")
            many_side = folder / "many.csv"
            many_side.write_text("ForeignKey\nmissing\n", encoding="utf-8")
            invalid_model = {
                "action": "model_audit",
                "tables": {"One": {"file": str(one_side)}, "Many": {"file": str(many_side)}},
                "relationships": [{"one_table": "One", "one_column": "Key", "many_table": "Many", "many_column": "ForeignKey"}],
            }
            report, summary = bi_report.build_report_result({
                "analyses": [
                    {"name": "Valid", "request": {"action": "inventory", "tables": TABLES}},
                    {"name": "Invalid audit", "request": invalid_model},
                    {"name": "Execution error", "request": {"action": "unsupported"}},
                    {"name": "Malformed"},
                ]
            })
            empty_report, empty_summary = bi_report.build_report_result({"analyses": []})
            failed_report, failed_summary = bi_report.build_report_result({
                "analyses": [
                    {"name": "Invalid audit", "request": invalid_model},
                    {"name": "Execution error", "request": {"action": "unsupported"}},
                ]
            })

        self.assertEqual(summary, {
            "status": "partial", "total": 4, "succeeded": 1, "invalid": 1, "errors": 1, "skipped": 1,
        })
        self.assertIn("REPORT STATUS: PARTIAL", report)
        self.assertEqual(empty_summary, {
            "status": "empty", "total": 0, "succeeded": 0, "invalid": 0, "errors": 0, "skipped": 0,
        })
        self.assertIn("REPORT STATUS: EMPTY", empty_report)
        self.assertEqual(failed_summary, {
            "status": "failed", "total": 2, "succeeded": 0, "invalid": 1, "errors": 1, "skipped": 0,
        })
        self.assertIn("REPORT STATUS: FAILED", failed_report)

    def test_report_cli_strict_mode_accepts_complete_and_rejects_empty(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            for label, analyses, expected_code, expected_status in (
                ("complete", [{"name": "Inventory", "request": {"action": "inventory", "tables": TABLES}}], 0, "COMPLETE"),
                ("empty", [], 2, "EMPTY"),
            ):
                config = folder / f"{label}.json"
                config.write_text(json.dumps({"analyses": analyses}), encoding="utf-8")
                output = folder / f"{label}.md"
                completed = subprocess.run(
                    [sys.executable, str(ROOT / "scripts" / "bi_report.py"), "--config", str(config),
                     "--output", str(output), "--fail-on-invalid"],
                    capture_output=True,
                    text=True,
                    check=False,
                )
                self.assertEqual(completed.returncode, expected_code, completed.stderr)
                self.assertIn(f"REPORT STATUS: {expected_status}", output.read_text(encoding="utf-8"))

    def test_report_builder_refuses_to_overwrite_its_request(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            config = folder / "report.json"
            config.write_text(json.dumps({"analyses": []}), encoding="utf-8")
            original = config.read_bytes()
            with redirect_stderr(io.StringIO()):
                status = bi_report.main(["--config", str(config), "--output", str(config)])
            self.assertEqual(status, 2)
            self.assertEqual(config.read_bytes(), original)

    def test_division_uses_documented_50_digit_precision(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "amounts.csv"
            path.write_text("Amount\n1\n2\n", encoding="utf-8")
            result = analyze.run_request({
                "action": "aggregate",
                "tables": {"Sales": {"file": str(path)}},
                "fact_table": "Sales",
                "measures": [{"name": "Mean", "op": "mean", "table": "Sales", "column": "Amount"}],
            })
        self.assertEqual(
            result["totals"]["Mean"],
            Decimal("1.5"),
        )
        self.assertEqual(len(analyze._decimal_divide(Decimal(1), Decimal(3)).as_tuple().digits), 50)

    def test_large_integral_decimal_serializes_without_python_int_limit(self):
        digits = "9" * 5_000
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "large.csv"
            path.write_text(f"Amount\n{digits}\n", encoding="utf-8")
            result = analyze.run_request({
                "action": "aggregate",
                "tables": {"Sales": {"file": str(path)}},
                "fact_table": "Sales",
                "measures": [{"name": "Total", "op": "sum", "table": "Sales", "column": "Amount"}],
            })
        serialized = analyze._json_dumps(result)
        self.assertIn(f'"Total": {digits}', serialized)
        decoded = json.loads(serialized, parse_int=Decimal, parse_float=Decimal)
        self.assertEqual(decoded["totals"]["Total"], Decimal(digits))

    def test_dimension_filter_is_applied_to_fact_rows(self):
        request = {
            "action": "aggregate",
            "tables": TABLES,
            "relationships": RELATIONSHIPS,
            "fact_table": "Orders",
            "filters": [{"table": "Customers", "column": "Region", "op": "eq", "value": "West"}],
            "measures": [{"name": "Net Sales", "op": "sum_expression", "expression": "Quantity * UnitPrice * (1 - DiscountPct)"}],
        }
        result = analyze.run_request(request)
        self.assertTrue(result["valid"])
        self.assertEqual(result["rows_used"], 5)
        self.assertEqual(result["totals"]["Net Sales"], 720)

    def test_arithmetic_expression_resource_limits(self):
        with patch.object(analyze.ast, "parse") as parse:
            with self.assertRaisesRegex(analyze.RequestError, "maximum length") as length_error:
                analyze._compile_expression("1" * (analyze._MAX_EXPRESSION_LENGTH + 1), {"Amount"})
        self.assertEqual(length_error.exception.code, "resource_limit")
        parse.assert_not_called()
        request = {
            "action": "aggregate",
            "tables": {"Sales": {"file": "/path/that/must/not/be/opened.csv"}},
            "fact_table": "Sales",
            "measures": [{
                "name": "Too long",
                "op": "sum_expression",
                "expression": "1" * (analyze._MAX_EXPRESSION_LENGTH + 1),
            }],
        }
        with self.assertRaisesRegex(analyze.RequestError, "maximum length") as preflight_error:
            analyze.run_request(request)
        self.assertEqual(preflight_error.exception.code, "resource_limit")
        self.assertEqual(preflight_error.exception.field_path, "measures[0].expression")

        def balanced_sum(term_count):
            if term_count == 1:
                return "1"
            midpoint = term_count // 2
            return f"({balanced_sum(midpoint)} + {balanced_sum(term_count - midpoint)})"

        many_nodes = balanced_sum(400)
        self.assertLess(len(many_nodes), analyze._MAX_EXPRESSION_LENGTH)
        with self.assertRaisesRegex(analyze.RequestError, "maximum AST node count") as node_error:
            analyze._compile_expression(many_nodes, {"Amount"})
        self.assertEqual(node_error.exception.code, "resource_limit")

        deep_expression = "-" * (analyze._MAX_EXPRESSION_AST_DEPTH + 1) + "1"
        with self.assertRaisesRegex(analyze.RequestError, "maximum AST depth") as depth_error:
            analyze._compile_expression(deep_expression, {"Amount"})
        self.assertEqual(depth_error.exception.code, "resource_limit")

    def test_expression_cannot_call_functions(self):
        with self.assertRaises(analyze.RequestError):
            analyze._compile_expression("__import__('os').system('echo unsafe')", {"Quantity"})

    def test_boolean_values_are_not_coerced_to_numeric_filters_or_measures(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            numeric_path = folder / "numeric.csv"
            numeric_path.write_text("Amount\n1\n0\n", encoding="utf-8")
            with self.assertRaisesRegex(analyze.RequestError, "Boolean value is not numeric"):
                analyze.run_request({
                    "action": "aggregate",
                    "tables": {"Sales": {"file": str(numeric_path)}},
                    "fact_table": "Sales",
                    "filters": [{"table": "Sales", "column": "Amount", "op": "eq", "value": True}],
                    "measures": [{"name": "Sum", "op": "sum", "table": "Sales", "column": "Amount"}],
                })

            import openpyxl
            boolean_path = folder / "booleans.xlsx"
            workbook = openpyxl.Workbook()
            worksheet = workbook.active
            worksheet.append(["Flag"])
            worksheet.append([True])
            worksheet.append([False])
            workbook.save(boolean_path)
            with self.assertRaisesRegex(analyze.RequestError, "Boolean value is not numeric"):
                analyze.run_request({
                    "action": "aggregate",
                    "tables": {"Flags": {"file": str(boolean_path)}},
                    "fact_table": "Flags",
                    "measures": [{"name": "Sum", "op": "sum", "table": "Flags", "column": "Flag"}],
                })

    def test_xlsx_archive_expansion_limit_precedes_workbook_loading(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "oversized.xlsx"
            with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("xl/worksheet.xml", b"x" * 101)
            with patch.object(analyze, "_MAX_XLSX_UNCOMPRESSED_BYTES", 100), patch("openpyxl.load_workbook") as load_workbook:
                with self.assertRaisesRegex(analyze.RequestError, "uncompressed.*maximum") as error:
                    analyze.load_table("Sales", {"file": str(path)})
        self.assertEqual(error.exception.code, "resource_limit")
        load_workbook.assert_not_called()

    def test_configurable_source_byte_limit_precedes_open_and_accepts_boundary(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "source.csv"
            source = b"Key\nvalue\n"
            path.write_bytes(source)
            request = {
                "action": "inventory",
                "tables": {"Sales": {"file": str(path)}},
                "resource_limits": {"max_source_bytes": len(source) - 1},
            }
            with patch.object(Path, "open", side_effect=AssertionError("source opened before size rejection")):
                with self.assertRaisesRegex(analyze.RequestError, "bytes.*maximum") as error:
                    analyze.run_request(request)
            self.assertEqual(error.exception.code, "resource_limit")
            self.assertEqual(error.exception.field_path, "tables.Sales.file")

            request["resource_limits"]["max_source_bytes"] = len(source)
            result = analyze.run_request(request)
            self.assertEqual(result["tables"]["Sales"]["rows"], 1)

            request.pop("resource_limits")
            with patch.object(analyze, "_DEFAULT_MAX_SOURCE_BYTES", len(source) - 1):
                with patch.object(Path, "open", side_effect=AssertionError("source opened before default-limit rejection")):
                    with self.assertRaises(analyze.RequestError) as default_error:
                        analyze.run_request(request)
            self.assertEqual(default_error.exception.code, "resource_limit")

    def test_invalid_source_limit_is_rejected_before_input_lookup(self):
        request = {
            "action": "inventory",
            "tables": {"Sales": {"file": "/path/that/must/not/be/opened.csv"}},
            "resource_limits": {"max_source_bytes": True},
        }
        with self.assertRaisesRegex(analyze.RequestError, "max_source_bytes.*integer") as error:
            analyze.run_request(request)
        self.assertEqual(error.exception.code, "invalid_type")
        self.assertEqual(error.exception.field_path, "resource_limits.max_source_bytes")

        request["resource_limits"]["max_source_bytes"] = analyze._MAX_CONFIGURABLE_SOURCE_BYTES + 1
        with self.assertRaisesRegex(analyze.RequestError, "max_source_bytes.*from 1") as out_of_range:
            analyze.run_request(request)
        self.assertEqual(out_of_range.exception.code, "invalid_value")
        self.assertEqual(out_of_range.exception.field_path, "resource_limits.max_source_bytes")

        request["resource_limits"] = {"max_rows": True}
        with self.assertRaisesRegex(analyze.RequestError, "max_rows.*integer") as row_type_error:
            analyze.run_request(request)
        self.assertEqual(row_type_error.exception.code, "invalid_type")
        self.assertEqual(row_type_error.exception.field_path, "resource_limits.max_rows")

        request["resource_limits"]["max_rows"] = 0
        with self.assertRaisesRegex(analyze.RequestError, "max_rows.*from 1") as row_range_error:
            analyze.run_request(request)
        self.assertEqual(row_range_error.exception.code, "invalid_value")

        request["resource_limits"] = {"max_groups": 100}
        with self.assertRaisesRegex(analyze.RequestError, "Unknown resource limit.*max_groups") as unknown:
            analyze.run_request(request)
        self.assertEqual(unknown.exception.code, "unknown_field")
        self.assertEqual(unknown.exception.field_path, "resource_limits.max_groups")

    def test_configurable_row_limit_rejects_csv_overflow_and_accepts_boundary(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "rows.csv"
            path.write_text("Key\nA\nB\nC\n", encoding="utf-8")
            request = {
                "action": "inventory",
                "tables": {"Sales": {"file": str(path)}},
                "resource_limits": {"max_rows": 2},
            }
            with self.assertRaisesRegex(analyze.RequestError, "maximum row count") as error:
                analyze.run_request(request)
            self.assertEqual(error.exception.code, "resource_limit")
            self.assertEqual(error.exception.field_path, "tables.Sales.file")

            config_path = Path(temp) / "request.json"
            config_path.write_text(json.dumps(request), encoding="utf-8")
            cli = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "analyze.py"), "--config", str(config_path)],
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(cli.returncode, 2)
            cli_result = json.loads(cli.stdout)
            self.assertEqual(cli_result["error_code"], "resource_limit")
            self.assertEqual(cli_result["error_path"], "tables.Sales.file")

            request["resource_limits"]["max_rows"] = 3
            result = analyze.run_request(request)
            self.assertEqual(result["tables"]["Sales"]["rows"], 3)

            request.pop("resource_limits")
            with patch.object(analyze, "_DEFAULT_MAX_ROWS", 2):
                with self.assertRaisesRegex(analyze.RequestError, "maximum row count") as default_error:
                    analyze.run_request(request)
            self.assertEqual(default_error.exception.code, "resource_limit")

    def test_xlsx_row_limit_consumes_rows_lazily(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "rows.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["Key"])
            sheet.append(["A"])
            sheet.append(["B"])
            workbook.save(path)
            workbook.close()

            with patch.object(analyze, "_headers_and_rows", wraps=analyze._headers_and_rows) as convert_rows:
                with self.assertRaisesRegex(analyze.RequestError, "maximum row count") as error:
                    analyze.load_table("Sales", {"file": str(path)}, max_rows=1)
            self.assertEqual(error.exception.code, "resource_limit")
            self.assertFalse(isinstance(convert_rows.call_args.args[1], list))

    def test_corrupt_xlsx_becomes_clean_cli_and_report_error(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            bad = folder / "bad.xlsx"
            bad.write_text("not an xlsx archive", encoding="utf-8")
            request = {"action": "inventory", "tables": {"Sales": {"file": str(bad)}}}
            config = folder / "request.json"
            config.write_text(json.dumps(request), encoding="utf-8")
            direct = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "analyze.py"), "--config", str(config)],
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(direct.returncode, 2)
            self.assertIn("Could not read Excel workbook", direct.stderr)
            self.assertNotIn("Traceback", direct.stderr)

            report_config = folder / "report.json"
            report_config.write_text(
                json.dumps({"analyses": [{"name": "Bad workbook", "request": request}]}),
                encoding="utf-8",
            )
            output = folder / "report.md"
            report_run = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "bi_report.py"), "--config", str(report_config), "--output", str(output)],
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(report_run.returncode, 0, report_run.stderr)
            report = output.read_text(encoding="utf-8")
            self.assertIn("Could not read Excel workbook", report)
            self.assertIn("## Limitations", report)

    def test_xlsx_named_sheet_is_read(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "sales.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Data"
            sheet.append(["Code", "Amount"])
            sheet.append(["A", 12.5])
            workbook.save(path)
            result = analyze.run_request({"action": "inventory", "tables": {"Sales": {"file": str(path), "sheet": "Data"}}})
        self.assertEqual(result["tables"]["Sales"]["rows"], 1)
        self.assertEqual(result["tables"]["Sales"]["fields"][1]["inferred_type"], "decimal")

    def test_header_only_csv_keeps_its_schema(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "empty.csv"
            path.write_text("Key,Value\n", encoding="utf-8")
            result = analyze.run_request({"action": "inventory", "tables": {"Empty": {"file": str(path)}}})
        self.assertEqual(result["tables"]["Empty"]["rows"], 0)
        self.assertEqual([field["name"] for field in result["tables"]["Empty"]["fields"]], ["Key", "Value"])

    def test_unicode_and_space_in_path(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp) / "BI data Δ"
            folder.mkdir()
            path = folder / "地域.csv"
            path.write_text("City,Amount\nMünchen,12\n", encoding="utf-8")
            result = analyze.run_request({"action": "inventory", "tables": {"Sales": {"file": str(path)}}, "include_samples": True})
        self.assertEqual(result["tables"]["Sales"]["fields"][0]["sample_values"], ["München"])

    def test_unknown_filter_column_is_a_request_error(self):
        request = {
            "action": "aggregate", "tables": TABLES, "relationships": RELATIONSHIPS,
            "fact_table": "Orders", "filters": [{"table": "Customers", "column": "NotAColumn", "value": "West"}],
            "measures": [{"name": "Rows", "op": "count_rows"}],
        }
        with self.assertRaises(analyze.RequestError):
            analyze.run_request(request)

    def test_non_string_filter_operations_are_request_errors(self):
        for operation in ([], {}):
            request = {
                "action": "aggregate", "tables": TABLES, "relationships": RELATIONSHIPS,
                "fact_table": "Orders",
                "filters": [{"table": "Orders", "column": "Quantity", "op": operation, "value": 1}],
                "measures": [{"name": "Rows", "op": "count_rows"}],
            }
            with self.subTest(operation=operation):
                with self.assertRaisesRegex(analyze.RequestError, r"filters\[0\]\.op must be a string"):
                    analyze.run_request(request)

    def test_non_string_measure_operations_are_request_errors(self):
        for operation in ([], {}):
            request = {
                "action": "aggregate", "tables": TABLES, "relationships": RELATIONSHIPS,
                "fact_table": "Orders", "measures": [{"name": "Rows", "op": operation}],
            }
            with self.subTest(operation=operation):
                with self.assertRaisesRegex(analyze.RequestError, r"measures\[0\]\.op must be a string"):
                    analyze.run_request(request)

    def test_nested_aggregate_contract_errors_are_rejected_before_file_io(self):
        tables = {"Sales": {"file": "/path/that/must/not/be/opened.csv"}}
        invalid_requests = (
            ({"filters": [{"table": "Sales", "column": "Amount", "op": [], "value": 1}]}, r"filters\[0\]\.op must be a string"),
            ({"filters": [{"table": "Sales", "column": "Amount", "op": {}, "value": 1}]}, r"filters\[0\]\.op must be a string"),
            ({"filters": [{"table": "Sales", "column": "Amount", "op": "contains", "value": 1}]}, r"Unsupported filters\[0\]\.op"),
            ({"filters": [{"table": "Sales", "column": "Amount", "op": "eq"}]}, r"filters\[0\]\.value is required"),
            ({"filters": [{"table": "Sales", "column": "Amount", "op": "in", "value": None}]}, r"filters\[0\]\.value must be a list"),
            ({"measures": [{"name": "Total", "op": []}]}, r"measures\[0\]\.op must be a string"),
            ({"measures": [{"name": "Total", "op": {}}]}, r"measures\[0\]\.op must be a string"),
            ({"measures": [{"name": "Total", "op": "median"}]}, r"Unsupported measures\[0\]\.op"),
        )
        for extra, expected in invalid_requests:
            request = {"action": "aggregate", "tables": tables, "fact_table": "Sales", "measures": [{"name": "Rows", "op": "count_rows"}], **extra}
            with self.subTest(extra=extra):
                with self.assertRaisesRegex(analyze.RequestError, expected):
                    analyze.run_request(request)

    def test_aggregate_field_shapes_and_duplicate_names_are_rejected_before_file_io(self):
        invalid_requests = (
            ({"group_by": {}}, r"'group_by' must be a list"),
            ({"group_by": [{"table": [], "column": "Amount"}]}, r"group_by\[0\].*string"),
            ({"group_by": [{"table": " ", "column": "Amount"}]}, r"group_by\[0\]\.table must not be empty"),
            ({"filters": [{"table": "Sales", "column": {}, "op": "eq", "value": 1}]}, r"filters\[0\].*string"),
            ({"measures": [{"name": "Total", "op": "sum", "table": "Sales"}]}, r"measures\[0\]\.column is required"),
            ({"measures": [{"name": "", "op": "count_rows"}]}, r"measures\[0\]\.name"),
            ({"measures": [{"name": "Rows", "op": "count_rows"}, {"name": "Rows", "op": "count_rows"}]}, r"Duplicate measure name at measures\[1\]\.name"),
            ({"measures": [{"name": "Expression", "op": "sum_expression", "expression": []}]}, r"measures\[0\]\.expression"),
        )
        with tempfile.TemporaryDirectory() as temp:
            tables = {"Sales": {"file": str(Path(temp) / "must-not-be-opened.csv")}}
            for extra, expected in invalid_requests:
                request = {"action": "aggregate", "tables": tables, "fact_table": "Sales", "measures": [{"name": "Rows", "op": "count_rows"}], **extra}
                with self.subTest(extra=extra):
                    with self.assertRaisesRegex(analyze.RequestError, expected):
                        analyze.run_request(request)

    def test_duplicate_measure_name_has_stable_error_metadata(self):
        request = {
            "action": "aggregate",
            "tables": {"Sales": {"file": "/path/that/must/not/be/opened.csv"}},
            "fact_table": "Sales",
            "measures": [
                {"name": "Rows", "op": "count_rows"},
                {"name": "Rows", "op": "count_rows"},
            ],
        }
        with self.assertRaises(analyze.RequestError) as context:
            analyze.run_request(request)
        self.assertEqual(context.exception.code, "duplicate_name")
        self.assertEqual(context.exception.field_path, "measures[1].name")

    def test_explicit_null_filter_value_matches_blank_values(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "sales.csv"
            source.write_text("Region,Amount\n,1\nWest,2\n", encoding="utf-8")
            base = {
                "action": "aggregate",
                "tables": {"Sales": {"file": str(source)}},
                "fact_table": "Sales",
                "measures": [{"name": "Rows", "op": "count_rows"}],
            }
            missing = {**base, "filters": [{"table": "Sales", "column": "Region", "op": "eq"}]}
            with self.assertRaisesRegex(analyze.RequestError, r"filters\[0\]\.value is required"):
                analyze.run_request(missing)
            explicit_null = {**base, "filters": [{"table": "Sales", "column": "Region", "op": "eq", "value": None}]}
            result = analyze.run_request(explicit_null)
        self.assertTrue(result["valid"])
        self.assertEqual(result["totals"]["Rows"], 1)

    def test_non_string_measure_operation_cli_exits_without_traceback(self):
        with tempfile.TemporaryDirectory() as temp:
            config_path = Path(temp) / "request.json"
            config_path.write_text(json.dumps({
                "action": "aggregate", "tables": TABLES, "relationships": RELATIONSHIPS,
                "fact_table": "Orders", "measures": [{"name": "Rows", "op": {}}],
            }), encoding="utf-8")
            process = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "analyze.py"), "--config", str(config_path)],
                capture_output=True, text=True, check=False,
            )
        self.assertEqual(process.returncode, 2)
        self.assertIn("measures[0].op must be a string", process.stderr)
        self.assertNotIn("Traceback", process.stderr)
        payload = json.loads(process.stdout)
        self.assertEqual(payload["error_code"], "invalid_type")
        self.assertEqual(payload["error_path"], "measures[0].op")

    def test_unsupported_measure_operation_cli_includes_stable_code_and_path(self):
        with tempfile.TemporaryDirectory() as temp:
            config_path = Path(temp) / "request.json"
            config_path.write_text(json.dumps({
                "action": "aggregate", "tables": TABLES, "relationships": RELATIONSHIPS,
                "fact_table": "Orders", "measures": [{"name": "Rows", "op": "median"}],
            }), encoding="utf-8")
            process = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "analyze.py"), "--config", str(config_path)],
                capture_output=True, text=True, check=False,
            )
        self.assertEqual(process.returncode, 2)
        self.assertNotIn("Traceback", process.stderr)
        self.assertTrue(process.stderr.startswith("Error: Unsupported measures[0].op"))
        payload = json.loads(process.stdout)
        self.assertEqual(payload["error_code"], "unsupported_operation")
        self.assertEqual(payload["error_path"], "measures[0].op")

    def test_unknown_top_level_filter_typo_is_rejected_before_file_io(self):
        request = {
            "action": "aggregate",
            "tables": {"Sales": {"file": "/path/that/must/not/be/opened.csv"}},
            "fact_table": "Sales",
            "filter": [{"table": "Sales", "column": "Region", "op": "eq", "value": "West"}],
            "measures": [{"name": "Rows", "op": "count_rows"}],
        }
        with self.assertRaisesRegex(
            analyze.RequestError,
            r"Unknown aggregate request key 'filter'; did you mean 'filters'\?",
        ):
            analyze.run_request(request)

    def test_non_string_action_is_rejected_before_file_io(self):
        request = {
            "action": [],
            "tables": {"Sales": {"file": "/path/that/must/not/be/opened.csv"}},
        }
        with self.assertRaisesRegex(analyze.RequestError, "'action' must be a string"):
            analyze.run_request(request)

    def test_output_cannot_overwrite_input(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "source.csv"
            config = {"tables": {"Sales": {"file": str(source)}}}
            with self.assertRaises(analyze.RequestError):
                analyze.ensure_output_safe(source, Path(temp) / "request.json", config)

    def test_output_cannot_overwrite_hard_link_to_source_or_request(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            source = folder / "source.csv"
            source.write_text("Amount\nPRIVATE-SENTINEL\n", encoding="utf-8")
            request_path = folder / "request.json"
            request_path.write_text('{"synthetic":"REQUEST-SENTINEL"}\n', encoding="utf-8")
            config = {"tables": {"Sales": {"file": str(source)}}}

            for index, reserved_path in enumerate((source, request_path)):
                alias = folder / f"output-alias-{index}"
                try:
                    os.link(reserved_path, alias)
                except OSError as exc:
                    self.skipTest(f"filesystem does not support hard links: {exc}")
                before = reserved_path.read_bytes()
                self.assertNotEqual(alias.resolve(), reserved_path.resolve())
                self.assertTrue(os.path.samefile(alias, reserved_path))
                with self.assertRaisesRegex(analyze.RequestError, "same file"):
                    analyze.ensure_output_safe(alias, request_path, config)
                self.assertEqual(reserved_path.read_bytes(), before)

    def test_both_cli_writers_reject_hard_link_to_source_without_modifying_it(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            source = folder / "source.csv"
            source.write_text("Amount\nPRIVATE-SENTINEL\n", encoding="utf-8")
            before = source.read_bytes()
            analysis = {"action": "inventory", "tables": {"Sales": {"file": str(source)}}}
            for script, config_object, suffix in (
                ("analyze.py", analysis, "analyze"),
                ("bi_report.py", {"analyses": [{"name": "Source", "request": analysis}]}, "report"),
            ):
                config_path = folder / f"{suffix}.json"
                config_path.write_text(json.dumps(config_object), encoding="utf-8")
                output = folder / f"{suffix}-output.md"
                try:
                    os.link(source, output)
                except OSError as exc:
                    self.skipTest(f"filesystem does not support hard links: {exc}")
                completed = subprocess.run(
                    [sys.executable, str(ROOT / "scripts" / script), "--config", str(config_path), "--output", str(output)],
                    capture_output=True,
                    text=True,
                    check=False,
                )
                self.assertEqual(completed.returncode, 2, completed.stderr)
                self.assertIn("same file", completed.stderr)
                self.assertEqual(source.read_bytes(), before)
                output.unlink()

    def test_cli_output_write_error_is_clean_and_nonzero(self):
        with tempfile.TemporaryDirectory() as temp:
            folder = Path(temp)
            config = folder / "request.json"
            config.write_text(json.dumps({
                "action": "inventory",
                "tables": {"Customers": {"file": str(DEMO / "customers.csv")}},
            }), encoding="utf-8")
            target = folder / "existing-directory"
            target.mkdir()
            completed = subprocess.run(
                [sys.executable, str(ROOT / "scripts" / "analyze.py"), "--config", str(config), "--output", str(target)],
                capture_output=True,
                text=True,
                check=False,
            )
        self.assertEqual(completed.returncode, 2)
        self.assertIn("Error writing output", completed.stderr)
        self.assertNotIn("Traceback", completed.stderr)

    def test_markdown_is_explicitly_labeled(self):
        result = analyze.run_request({"action": "inventory", "tables": TABLES})
        rendered = analyze.format_markdown(result)
        self.assertIn("not Power BI rendering or DAX evaluation", rendered)
        self.assertIn("## Orders", rendered)


if __name__ == "__main__":
    unittest.main()
