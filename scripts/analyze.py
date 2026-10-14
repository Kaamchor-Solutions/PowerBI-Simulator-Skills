#!/usr/bin/env python3
"""Small, transparent Power BI-style data helper; this is not a DAX engine."""
from __future__ import annotations

import argparse
import ast
import csv
import datetime as dt
import difflib
import json
import math
import sys
from collections import Counter, defaultdict
from decimal import (
    Clamped,
    Decimal,
    DecimalException,
    DivisionByZero,
    Inexact,
    InvalidOperation,
    MAX_EMAX,
    MIN_EMIN,
    Overflow,
    Rounded,
    ROUND_HALF_EVEN,
    Subnormal,
    Underflow,
    localcontext,
)
from pathlib import Path
from typing import Any, Callable


class RequestError(ValueError):
    """A request or source file is invalid or outside the helper's scope."""

    def __init__(self, message: str, *, code: str = "invalid_request", field_path: str | None = None):
        super().__init__(message)
        self.code = code
        self.field_path = field_path


_SAMPLE_POLICY_KEYS = frozenset({"include_samples", "sample_limit"})
_REQUEST_KEYS = {
    "inventory": frozenset({"action", "tables"}) | _SAMPLE_POLICY_KEYS,
    "model_audit": frozenset({"action", "tables", "relationships"}) | _SAMPLE_POLICY_KEYS,
    "aggregate": frozenset({
        "action", "tables", "relationships", "fact_table", "group_by", "filters", "measures",
    }) | _SAMPLE_POLICY_KEYS,
}
_FILTER_OPERATIONS = frozenset({"eq", "neq", "in"})
_MEASURE_OPERATIONS = frozenset({
    "count_rows", "count_nonblank", "distinct_count", "sum", "mean", "min", "max", "sum_expression",
})
_MAX_EXPRESSION_LENGTH = 4096
_MAX_EXPRESSION_AST_NODES = 1024
_MAX_EXPRESSION_AST_DEPTH = 64


def _validate_request_keys(config: dict[str, Any]) -> str:
    action = config.get("action")
    if not isinstance(action, str):
        raise RequestError("'action' must be a string: inventory, model_audit, aggregate", code="invalid_type", field_path="action")
    if action not in _REQUEST_KEYS:
        raise RequestError("'action' must be one of: inventory, model_audit, aggregate", code="unsupported_value", field_path="action")

    allowed = _REQUEST_KEYS[action]
    unknown = sorted((key for key in config if key not in allowed), key=str)
    if unknown:
        key = unknown[0]
        message = f"Unknown {action} request key {key!r}"
        suggestion = difflib.get_close_matches(str(key), sorted(allowed), n=1, cutoff=0.6)
        if suggestion:
            message += f"; did you mean {suggestion[0]!r}?"
        raise RequestError(message, code="unknown_field", field_path=str(key))
    return action


def _sample_policy(config: dict[str, Any]) -> tuple[bool, int]:
    include_samples = config.get("include_samples", False)
    if not isinstance(include_samples, bool):
        raise RequestError("'include_samples' must be a boolean", code="invalid_type", field_path="include_samples")
    sample_limit = config.get("sample_limit", 5)
    if isinstance(sample_limit, bool) or not isinstance(sample_limit, int) or not 0 <= sample_limit <= 5:
        raise RequestError("'sample_limit' must be an integer from 0 to 5", code="invalid_value", field_path="sample_limit")
    return include_samples, sample_limit


def _validate_aggregate_operations(config: dict[str, Any]) -> None:
    group_by = config.get("group_by", [])
    filters = config.get("filters", [])
    if not isinstance(group_by, list):
        raise RequestError("'group_by' must be a list", code="invalid_type", field_path="group_by")
    if not isinstance(filters, list):
        raise RequestError("'filters' must be a list", code="invalid_type", field_path="filters")
    if "measures" not in config:
        raise RequestError("'measures' is required", code="missing_field", field_path="measures")
    measures = config["measures"]
    if not isinstance(measures, list):
        raise RequestError("'measures' must be a list", code="invalid_type", field_path="measures")
    if not measures:
        raise RequestError("'measures' must be a non-empty list", code="invalid_value", field_path="measures")
    for index, item in enumerate(group_by):
        _field_spec(item, f"group_by[{index}]")
    for index, item in enumerate(filters):
        if not isinstance(item, dict):
            raise RequestError(f"filters[{index}] must be an object", code="invalid_type", field_path=f"filters[{index}]")
        _field_spec(item, f"filters[{index}]")
        operation = item.get("op", "eq")
        if not isinstance(operation, str):
            raise RequestError(f"filters[{index}].op must be a string", code="invalid_type", field_path=f"filters[{index}].op")
        if operation not in _FILTER_OPERATIONS:
            raise RequestError(f"Unsupported filters[{index}].op {operation!r}; use {sorted(_FILTER_OPERATIONS)}", code="unsupported_operation", field_path=f"filters[{index}].op")
        if "value" not in item:
            raise RequestError(f"filters[{index}].value is required; use null explicitly to match blank values", code="missing_field", field_path=f"filters[{index}].value")
        if operation == "in" and not isinstance(item.get("value"), list):
            raise RequestError(f"filters[{index}].value must be a list when op is 'in'", code="invalid_type", field_path=f"filters[{index}].value")
    measure_names: set[str] = set()
    for index, measure in enumerate(measures):
        if not isinstance(measure, dict):
            raise RequestError(f"measures[{index}] must be an object", code="invalid_type", field_path=f"measures[{index}]")
        if "name" not in measure:
            raise RequestError(f"measures[{index}].name is required", code="missing_field", field_path=f"measures[{index}].name")
        name = measure.get("name")
        if not isinstance(name, str):
            raise RequestError(f"measures[{index}].name must be a string", code="invalid_type", field_path=f"measures[{index}].name")
        if not name.strip():
            raise RequestError(f"measures[{index}].name must not be empty", code="invalid_value", field_path=f"measures[{index}].name")
        if name in measure_names:
            raise RequestError(f"Duplicate measure name at measures[{index}].name: {name!r}", code="duplicate_name", field_path=f"measures[{index}].name")
        measure_names.add(name)
        if "op" not in measure:
            raise RequestError(f"measures[{index}].op is required", code="missing_field", field_path=f"measures[{index}].op")
        operation = measure.get("op")
        if not isinstance(operation, str):
            raise RequestError(f"measures[{index}].op must be a string", code="invalid_type", field_path=f"measures[{index}].op")
        if operation not in _MEASURE_OPERATIONS:
            raise RequestError(f"Unsupported measures[{index}].op {operation!r}; use {sorted(_MEASURE_OPERATIONS)}", code="unsupported_operation", field_path=f"measures[{index}].op")
        if operation == "sum_expression":
            if "expression" not in measure:
                raise RequestError(f"measures[{index}].expression is required", code="missing_field", field_path=f"measures[{index}].expression")
            expression = measure.get("expression")
            if not isinstance(expression, str):
                raise RequestError(f"measures[{index}].expression must be a string", code="invalid_type", field_path=f"measures[{index}].expression")
            if not expression.strip():
                raise RequestError(f"measures[{index}].expression must not be empty", code="invalid_value", field_path=f"measures[{index}].expression")
            if len(expression) > _MAX_EXPRESSION_LENGTH:
                raise RequestError(
                    f"Arithmetic expression exceeds maximum length of {_MAX_EXPRESSION_LENGTH} characters",
                    code="resource_limit",
                    field_path=f"measures[{index}].expression",
                )
        elif operation != "count_rows":
            _field_spec(measure, f"measures[{index}]")


# Addition and multiplication of finite Decimal inputs are exact unless the
# active context is too small. Bound pathological exponent/alignment requests
# so an untrusted input cannot demand unbounded memory for one arithmetic step.
_MAX_EXACT_PRECISION = 100_000
_DIVISION_PRECISION = 50
_MAX_JSON_NUMBER_TOKEN_LENGTH = 100_000


def _reject_json_constant(token: str):
    raise RequestError(f"Non-standard JSON numeric constant {token!r} is not allowed", code="invalid_json_constant")


def _reject_duplicate_json_keys(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise RequestError(f"Duplicate JSON object key {key!r} is not allowed", code="duplicate_json_key")
        result[key] = value
    return result


def _parse_json_decimal(token: str) -> Decimal:
    if len(token) > _MAX_JSON_NUMBER_TOKEN_LENGTH:
        raise RequestError("JSON number token exceeds the 100,000-character limit")
    try:
        value = Decimal(token)
    except (InvalidOperation, ValueError) as exc:
        raise RequestError("JSON decimal number is outside the supported Decimal range") from exc
    if not value.is_finite():
        raise RequestError("JSON number must be finite")
    return value


def _parse_json_integer(token: str) -> int | Decimal:
    if len(token) > _MAX_JSON_NUMBER_TOKEN_LENGTH:
        raise RequestError("JSON number token exceeds the 100,000-character limit")
    try:
        return int(token)
    except ValueError:
        # Python limits conversion of very long integer strings. Preserve those
        # values as Decimal so numeric filters stay exact without a traceback.
        return _parse_json_decimal(token)


def load_json(handle) -> Any:
    """Load JSON without converting decimal number tokens through binary float."""
    return json.load(
        handle,
        parse_int=_parse_json_integer,
        parse_float=_parse_json_decimal,
        parse_constant=_reject_json_constant,
        object_pairs_hook=_reject_duplicate_json_keys,
    )


class TableRows(list[dict[str, Any]]):
    """Row records plus headers, including when the source contains no rows."""

    def __init__(self, rows: list[dict[str, Any]], columns: list[str]):
        super().__init__(rows)
        self.columns = list(columns)


def _clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, str) and not value.strip():
        return None
    return value


def _json_value(value: Any) -> Any:
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        if not value.is_finite():
            return None
        if value == value.to_integral_value():
            integer_digits = value.adjusted() + 1 if not value.is_zero() else 1
            get_int_limit = getattr(sys, "get_int_max_str_digits", None)
            int_string_limit = (get_int_limit() if get_int_limit is not None else 4_300) or 4_300
            if integer_digits <= int_string_limit:
                return int(value)
        return value
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _json_dumps(value: Any, *, ensure_ascii: bool = False, indent: int | None = None, sort_keys: bool = False) -> str:
    """Serialize JSON while preserving finite Decimal values as numeric tokens."""
    def encode(item: Any, level: int = 0) -> str:
        if isinstance(item, Decimal):
            return str(item) if item.is_finite() else "null"
        if isinstance(item, float) and not math.isfinite(item):
            return "null"
        if item is None or isinstance(item, (str, int, float, bool)):
            return json.dumps(item, ensure_ascii=ensure_ascii, allow_nan=False)
        if isinstance(item, (list, tuple)):
            if not item:
                return "[]"
            parts = [encode(child, level + 1) for child in item]
            if indent is None:
                return "[" + ", ".join(parts) + "]"
            pad = " " * (indent * (level + 1))
            close_pad = " " * (indent * level)
            return "[\n" + pad + (",\n" + pad).join(parts) + "\n" + close_pad + "]"
        if isinstance(item, dict):
            pairs = list(item.items())
            if sort_keys:
                pairs.sort(key=lambda pair: str(pair[0]))
            if not pairs:
                return "{}"
            if any(not isinstance(key, str) for key, _ in pairs):
                raise TypeError("JSON object keys must be strings")
            parts = [(json.dumps(key, ensure_ascii=ensure_ascii), encode(child, level + 1)) for key, child in pairs]
            if indent is None:
                return "{" + ", ".join(key + ": " + child for key, child in parts) + "}"
            pad = " " * (indent * (level + 1))
            close_pad = " " * (indent * level)
            return "{\n" + pad + (",\n" + pad).join(key + ": " + child for key, child in parts) + "\n" + close_pad + "}"
        raise TypeError(f"Object of type {type(item).__name__} is not JSON serializable")

    return encode(value)


def _headers_and_rows(headers: list[Any], data_rows: list[tuple[Any, ...]], table_name: str) -> TableRows:
    clean_headers = [str(h).strip() if h is not None else "" for h in headers]
    if not clean_headers or any(not name for name in clean_headers):
        raise RequestError(f"Table {table_name!r} needs a non-empty header row")
    if len(clean_headers) != len(set(clean_headers)):
        raise RequestError(f"Table {table_name!r} has duplicate column names")
    rows: list[dict[str, Any]] = []
    for row_number, values in enumerate(data_rows, start=2):
        if len(values) > len(clean_headers) and any(v is not None for v in values[len(clean_headers):]):
            raise RequestError(f"Table {table_name!r}, row {row_number} has more values than columns")
        padded = list(values[:len(clean_headers)]) + [None] * max(0, len(clean_headers) - len(values))
        row = {name: _clean_cell(value) for name, value in zip(clean_headers, padded)}
        if any(value is not None for value in row.values()):
            rows.append(row)
    return TableRows(rows, clean_headers)


def load_table(table_name: str, spec: dict[str, Any]) -> list[dict[str, Any]]:
    if not isinstance(spec, dict) or not isinstance(spec.get("file"), str):
        raise RequestError(f"Table {table_name!r} must define a string 'file' path")
    path = Path(spec["file"]).expanduser()
    if not path.is_file():
        raise RequestError(f"Input file for table {table_name!r} was not found: {path}")
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv", ".tab"}:
        delimiter = "\t" if suffix in {".tsv", ".tab"} else ","
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.reader(handle, delimiter=delimiter)
                try:
                    headers = next(reader)
                except StopIteration as exc:
                    raise RequestError(f"Input file is empty: {path}") from exc
                headers = [str(name).strip() for name in headers]
                if not headers or any(not str(h).strip() for h in headers):
                    raise RequestError(f"Input file has a blank header: {path}")
                if len(headers) != len(set(headers)):
                    raise RequestError(f"Input file has duplicate column names: {path}")
                rows = []
                for line_number, values in enumerate(reader, start=2):
                    if len(values) > len(headers):
                        raise RequestError(f"{path}, row {line_number} has more values than headers")
                    if not values or not any(value.strip() for value in values):
                        continue
                    padded = values + [""] * (len(headers) - len(values))
                    rows.append({name: _clean_cell(value) for name, value in zip(headers, padded)})
                return TableRows(rows, headers)
        except UnicodeDecodeError as exc:
            raise RequestError(f"Could not decode {path} as UTF-8; save as UTF-8 CSV/TSV") from exc
        except csv.Error as exc:
            raise RequestError(f"Could not parse {path}: {exc}") from exc
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RequestError("Reading .xlsx requires openpyxl; install scripts/requirements.txt") from exc
        workbook = None
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = spec.get("sheet", 0)
            if isinstance(sheet, str):
                if sheet not in workbook.sheetnames:
                    raise RequestError(f"Sheet {sheet!r} not found in {path}; available: {workbook.sheetnames}")
                worksheet = workbook[sheet]
            elif isinstance(sheet, int) and not isinstance(sheet, bool):
                if sheet < 0 or sheet >= len(workbook.worksheets):
                    raise RequestError(f"Sheet index {sheet} is out of range for {path}")
                worksheet = workbook.worksheets[sheet]
            else:
                raise RequestError("Excel 'sheet' must be a sheet name or zero-based integer index")
            iterator = worksheet.iter_rows(values_only=True)
            try:
                headers = list(next(iterator))
            except StopIteration as exc:
                raise RequestError(f"Excel sheet is empty: {path}") from exc
            return _headers_and_rows(headers, list(iterator), table_name)
        except RequestError:
            raise
        except Exception as exc:
            raise RequestError(f"Could not read Excel workbook {path}: {exc}") from exc
        finally:
            if workbook is not None:
                workbook.close()
    raise RequestError(f"Unsupported file for {table_name!r}: {suffix}. Supported: .csv, .tsv, .tab, .xlsx")


def load_tables(config: dict[str, Any]) -> dict[str, list[dict[str, Any]]]:
    specs = config.get("tables")
    if not isinstance(specs, dict) or not specs:
        raise RequestError("Request must contain a non-empty 'tables' object")
    tables: dict[str, list[dict[str, Any]]] = {}
    for name, spec in specs.items():
        if not isinstance(name, str) or not name.strip():
            raise RequestError("Table names must be non-empty strings")
        tables[name] = load_table(name, spec)
    return tables


def _decimal(value: Any, label: str) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, bool):
        raise RequestError(f"Boolean value is not numeric for {label}: {value!r}")
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError) as exc:
        raise RequestError(f"Expected a finite numeric value for {label}; found {value!r}") from exc
    if not number.is_finite():
        raise RequestError(f"Non-finite numeric value for {label}: {value!r}")
    return number


def _with_decimal_context(operation, precision: int) -> Decimal:
    if precision > _MAX_EXACT_PRECISION:
        raise RequestError(
            f"Exact arithmetic would require more than {_MAX_EXACT_PRECISION:,} digits"
        )
    try:
        with localcontext() as context:
            context.prec = max(1, precision)
            context.rounding = ROUND_HALF_EVEN
            context.clamp = 0
            context.Emax = MAX_EMAX
            context.Emin = MIN_EMIN
            # Pin relevant settings so arithmetic does not depend on the
            # caller's Decimal context; exact operations may still signal
            # Rounded/Subnormal without losing digits.
            for signal in (Clamped, Inexact, Rounded, Subnormal, Underflow):
                context.traps[signal] = False
            for signal in (DivisionByZero, InvalidOperation, Overflow):
                context.traps[signal] = True
            return operation()
    except DecimalException as exc:
        raise RequestError("Numeric result exceeds the supported Decimal range") from exc


def _decimal_add(left: Decimal, right: Decimal) -> Decimal:
    if left.is_zero():
        return right
    if right.is_zero():
        return left
    left_tuple, right_tuple = left.as_tuple(), right.as_tuple()
    base_exponent = min(left_tuple.exponent, right_tuple.exponent)
    precision = max(
        len(left_tuple.digits) + left_tuple.exponent - base_exponent,
        len(right_tuple.digits) + right_tuple.exponent - base_exponent,
    ) + 1
    return _with_decimal_context(lambda: left + right, precision)


def _decimal_subtract(left: Decimal, right: Decimal) -> Decimal:
    if right.is_zero():
        return left
    return _decimal_add(left, right.copy_negate())


def _decimal_multiply(left: Decimal, right: Decimal) -> Decimal:
    if left.is_zero() or right.is_zero():
        return Decimal(0)
    precision = len(left.as_tuple().digits) + len(right.as_tuple().digits)
    return _with_decimal_context(lambda: left * right, precision)


def _decimal_sum(values: list[Decimal]) -> Decimal:
    total = Decimal(0)
    for value in values:
        total = _decimal_add(total, value)
    return total


def _decimal_divide(numerator: Decimal, denominator: Decimal) -> Decimal:
    if denominator.is_zero():
        raise RequestError("Division by zero")
    return _with_decimal_context(
        lambda: numerator / denominator, _DIVISION_PRECISION
    )


def _parse_iso_date(value: Any) -> bool:
    if isinstance(value, (dt.date, dt.datetime)):
        return True
    if not isinstance(value, str):
        return False
    try:
        dt.date.fromisoformat(value[:10])
        return len(value) == 10 or (len(value) > 10 and value[10] in {"T", " "})
    except ValueError:
        return False


def infer_type(values: list[Any]) -> str:
    present = [value for value in values if value is not None]
    if not present:
        return "empty"
    if all(isinstance(value, bool) for value in present):
        return "boolean"
    if all(_parse_iso_date(value) for value in present):
        return "date/datetime"
    numeric: list[Decimal] = []
    for value in present:
        if isinstance(value, str) and len(value) > 1 and value.startswith("0") and value[1].isdigit() and "." not in value:
            return "text (leading-zero values preserved)"
        try:
            parsed = Decimal(str(value).strip())
        except (InvalidOperation, ValueError):
            return "text or mixed"
        if not parsed.is_finite():
            return "text or mixed"
        numeric.append(parsed)
    return "integer" if all(value == value.to_integral_value() for value in numeric) else "decimal"


def inventory(config: dict[str, Any], tables: dict[str, list[dict[str, Any]]]) -> dict[str, Any]:
    include_samples, sample_limit = _sample_policy(config)
    summaries: dict[str, Any] = {}
    for name, rows in tables.items():
        columns = list(rows.columns)
        duplicate_count = sum(count - 1 for count in Counter(
            _json_dumps({key: _json_value(value) for key, value in row.items()}, sort_keys=True)
            for row in rows
        ).values())
        fields = []
        for column in columns:
            values = [row.get(column) for row in rows]
            nonmissing = [value for value in values if value is not None]
            fields.append({
                "name": column,
                "inferred_type": infer_type(values),
                "nonmissing": len(nonmissing),
                "missing": len(values) - len(nonmissing),
                "distinct": len({_key(value) for value in nonmissing}),
                "sample_values": (
                    [_json_value(value) for value in nonmissing[:sample_limit]]
                    if include_samples else []
                ),
            })
        summaries[name] = {
            "file": config["tables"][name]["file"],
            "rows": len(rows),
            "columns": len(columns),
            "duplicate_rows": duplicate_count,
            "fields": fields,
        }
    return {"action": "inventory", "valid": True, "status": "valid", "tables": summaries}


def _key(value: Any) -> str:
    return _json_dumps(_json_value(value), sort_keys=True)


def _table_columns(tables: dict[str, list[dict[str, Any]]], table: str) -> list[str]:
    if table not in tables:
        raise RequestError(f"Unknown table {table!r}; available: {list(tables)}")
    rows = tables[table]
    return list(getattr(rows, "columns", list(rows[0].keys()) if rows else []))


def _column(tables: dict[str, list[dict[str, Any]]], table: str, column: str) -> list[Any]:
    if column not in _table_columns(tables, table):
        raise RequestError(f"Column {column!r} not found in table {table!r}")
    return [row.get(column) for row in tables[table]]


def audit_relationships(config: dict[str, Any], tables: dict[str, list[dict[str, Any]]]) -> dict[str, Any]:
    include_samples, sample_limit = _sample_policy(config)
    relationships = config.get("relationships", [])
    if not isinstance(relationships, list):
        raise RequestError("'relationships' must be a list")
    results = []
    for index, rel in enumerate(relationships):
        if not isinstance(rel, dict):
            raise RequestError(f"Relationship {index + 1} must be an object")
        required = ("one_table", "one_column", "many_table", "many_column")
        missing = [field for field in required if not isinstance(rel.get(field), str) or not rel[field]]
        if missing:
            raise RequestError(f"Relationship {index + 1} is missing string fields: {missing}")
        one_table, one_column = rel["one_table"], rel["one_column"]
        many_table, many_column = rel["many_table"], rel["many_column"]
        if one_table == many_table:
            raise RequestError(f"Relationship {index + 1} must connect two different tables")
        one_values = _column(tables, one_table, one_column)
        many_values = _column(tables, many_table, many_column)
        nonnull_one = [value for value in one_values if value is not None]
        counts = Counter(_key(value) for value in nonnull_one)
        key_samples: dict[str, Any] = {}
        for value in nonnull_one:
            key_samples.setdefault(_key(value), value)
        duplicate_keys = [key_samples[key] for key, count in counts.items() if count > 1]
        one_keys = set(counts)
        orphan_values = [value for value in many_values if value is not None and _key(value) not in one_keys]
        null_one = sum(value is None for value in one_values)
        null_many = sum(value is None for value in many_values)
        valid = not duplicate_keys and not orphan_values and null_one == 0
        results.append({
            "one_table": one_table,
            "one_column": one_column,
            "many_table": many_table,
            "many_column": many_column,
            "valid": valid,
            "one_key_nulls": null_one,
            "duplicate_one_key_count": len(duplicate_keys),
            "duplicate_one_key_samples": (
                [_json_value(value) for value in duplicate_keys[:sample_limit]]
                if include_samples else []
            ),
            "many_key_nulls": null_many,
            "orphan_count": len(orphan_values),
            "orphan_samples": (
                [_json_value(value) for value in list(dict.fromkeys(orphan_values))[:sample_limit]]
                if include_samples else []
            ),
        })
    valid = all(result["valid"] for result in results)
    warnings = [] if results else ["No relationships were declared; relationship behavior was not checked."]
    return {
        "action": "model_audit",
        "valid": valid,
        "status": "valid" if valid else "invalid",
        "relationships": results,
        "warnings": warnings,
        "scope_note": "Checks declared direct key uniqueness and referential integrity only; it does not implement Power BI relationship or filter propagation.",
    }


def _related_row(
    fact_row: dict[str, Any], field: dict[str, str], fact_table: str,
    tables: dict[str, list[dict[str, Any]]], relationships: list[dict[str, Any]],
    indexes: dict[tuple[str, str, str], dict[str, dict[str, Any]]],
) -> Any:
    table, column = field["table"], field["column"]
    _validate_field_relation(field, fact_table, tables, relationships)
    if table == fact_table:
        return fact_row.get(column)
    matches = [rel for rel in relationships if rel["many_table"] == fact_table and rel["one_table"] == table]
    if len(matches) != 1:
        raise RequestError(f"Field {table}.{column} needs exactly one declared direct relationship to fact table {fact_table!r}")
    rel = matches[0]
    fk = fact_row.get(rel["many_column"])
    lookup = indexes[(rel["one_table"], rel["one_column"], rel["many_table"])]
    related = lookup.get(_key(fk)) if fk is not None else None
    return related.get(column) if related else None


def _field_spec(value: Any, label: str) -> dict[str, str]:
    if not isinstance(value, dict):
        raise RequestError(f"{label} must be an object with non-empty string 'table' and 'column' fields", code="invalid_type", field_path=label)
    if "table" not in value:
        raise RequestError(f"{label}.table is required", code="missing_field", field_path=f"{label}.table")
    if "column" not in value:
        raise RequestError(f"{label}.column is required", code="missing_field", field_path=f"{label}.column")
    table, column = value.get("table"), value.get("column")
    if not isinstance(table, str):
        raise RequestError(f"{label}.table must be a string", code="invalid_type", field_path=f"{label}.table")
    if not table.strip():
        raise RequestError(f"{label}.table must not be empty", code="invalid_value", field_path=f"{label}.table")
    if not isinstance(column, str):
        raise RequestError(f"{label}.column must be a string", code="invalid_type", field_path=f"{label}.column")
    if not column.strip():
        raise RequestError(f"{label}.column must not be empty", code="invalid_value", field_path=f"{label}.column")
    return {"table": table, "column": column}


def _validate_field_relation(
    field: dict[str, str], fact_table: str, tables: dict[str, list[dict[str, Any]]],
    relationships: list[dict[str, Any]],
) -> None:
    table, column = field["table"], field["column"]
    if column not in _table_columns(tables, table):
        raise RequestError(f"Column {column!r} not found in table {table!r}")
    if table == fact_table:
        return
    matches = [rel for rel in relationships if rel["many_table"] == fact_table and rel["one_table"] == table]
    if len(matches) != 1:
        raise RequestError(f"Field {table}.{column} needs exactly one declared direct relationship to fact table {fact_table!r}")
    rel = matches[0]
    if rel["many_column"] not in _table_columns(tables, fact_table):
        raise RequestError(f"Column {rel['many_column']!r} not found in table {fact_table!r}")
    if rel["one_column"] not in _table_columns(tables, table):
        raise RequestError(f"Column {rel['one_column']!r} not found in table {table!r}")


def _compile_expression(expression: Any, fact_columns: set[str]):
    if not isinstance(expression, str) or not expression.strip():
        raise RequestError("sum_expression requires a non-empty arithmetic 'expression'")
    if len(expression) > _MAX_EXPRESSION_LENGTH:
        raise RequestError(
            f"Arithmetic expression exceeds maximum length of {_MAX_EXPRESSION_LENGTH} characters",
            code="resource_limit",
        )
    try:
        tree = ast.parse(expression, mode="eval")
    except RecursionError as exc:
        raise RequestError("Arithmetic expression exceeds the parser recursion limit", code="resource_limit") from exc
    except SyntaxError as exc:
        raise RequestError(f"Invalid arithmetic expression: {exc.msg}") from exc

    node_count = 0
    pending = [(tree, 1)]
    while pending:
        node, depth = pending.pop()
        node_count += 1
        if node_count > _MAX_EXPRESSION_AST_NODES:
            raise RequestError(
                f"Arithmetic expression exceeds maximum AST node count of {_MAX_EXPRESSION_AST_NODES}",
                code="resource_limit",
            )
        if depth > _MAX_EXPRESSION_AST_DEPTH:
            raise RequestError(
                f"Arithmetic expression exceeds maximum AST depth of {_MAX_EXPRESSION_AST_DEPTH}",
                code="resource_limit",
            )
        pending.extend((child, depth + 1) for child in ast.iter_child_nodes(node))

    allowed_binary = (ast.Add, ast.Sub, ast.Mult, ast.Div)
    literals: dict[int, Decimal] = {}

    def numeric_literal(node: ast.Constant) -> Decimal:
        if isinstance(node.value, int):
            return Decimal(node.value)
        source = ast.get_source_segment(expression, node)
        try:
            value = Decimal(source) if source is not None else Decimal(str(node.value))
        except (InvalidOperation, ValueError) as exc:
            raise RequestError("Invalid numeric literal in arithmetic expression") from exc
        if not value.is_finite():
            raise RequestError("Arithmetic expression numeric literals must be finite")
        return value

    def check(node: ast.AST) -> None:
        if isinstance(node, ast.Expression):
            check(node.body)
        elif isinstance(node, ast.BinOp) and isinstance(node.op, allowed_binary):
            check(node.left)
            check(node.right)
        elif isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            check(node.operand)
        elif isinstance(node, ast.Name):
            if node.id not in fact_columns:
                raise RequestError(f"Expression references unknown fact column {node.id!r}")
        elif isinstance(node, ast.Constant) and isinstance(node.value, (int, float)) and not isinstance(node.value, bool):
            literals[id(node)] = numeric_literal(node)
            return
        else:
            raise RequestError("Expressions may use only column names, numeric constants, +, -, *, /, and unary +/-; functions and code execution are not supported")

    check(tree)

    def evaluate_node(node: ast.AST, row: dict[str, Any]) -> Decimal | None:
        if isinstance(node, ast.Expression):
            return evaluate_node(node.body, row)
        if isinstance(node, ast.Constant):
            return literals[id(node)]
        if isinstance(node, ast.Name):
            return _decimal(row.get(node.id), f"expression column {node.id}")
        if isinstance(node, ast.UnaryOp):
            value = evaluate_node(node.operand, row)
            if value is None:
                return None
            return value if isinstance(node.op, ast.UAdd) else value.copy_negate()
        if isinstance(node, ast.BinOp):
            left, right = evaluate_node(node.left, row), evaluate_node(node.right, row)
            if left is None or right is None:
                return None
            if isinstance(node.op, ast.Add):
                return _decimal_add(left, right)
            if isinstance(node.op, ast.Sub):
                return _decimal_subtract(left, right)
            if isinstance(node.op, ast.Mult):
                return _decimal_multiply(left, right)
            if isinstance(node.op, ast.Div):
                if right == 0:
                    raise RequestError("Division by zero in sum_expression")
                return _decimal_divide(left, right)
        raise RequestError("Unsupported arithmetic expression")

    return lambda row: evaluate_node(tree, row)


def _filter_value(rows: list[dict[str, Any]], column: str, target: Any) -> Any:
    kind = infer_type([row.get(column) for row in rows])
    if kind in {"integer", "decimal"}:
        return _decimal(target, f"filter for {column}")
    return _clean_cell(target)


def _matches_filter(actual: Any, target: Any) -> bool:
    if actual is None:
        return target is None
    if isinstance(target, Decimal):
        actual_decimal = _decimal(actual, "filter column")
        return actual_decimal == target
    return _key(actual) == _key(target)


def _apply_filters(
    fact_rows: list[dict[str, Any]], filters: list[dict[str, Any]], fact_table: str,
    tables: dict[str, list[dict[str, Any]]], relationships: list[dict[str, Any]],
    indexes: dict[tuple[str, str, str], dict[str, dict[str, Any]]],
) -> list[dict[str, Any]]:
    prepared = []
    for index, item in enumerate(filters):
        field = _field_spec(item, f"Filter {index + 1}")
        _validate_field_relation(field, fact_table, tables, relationships)
        operation = item.get("op", "eq")
        if not isinstance(operation, str):
            raise RequestError(f"filters[{index}].op must be a string")
        if operation not in _FILTER_OPERATIONS:
            raise RequestError(f"Unsupported filters[{index}].op {operation!r}; use {sorted(_FILTER_OPERATIONS)}")
        values = item.get("value")
        if operation == "in":
            if not isinstance(values, list):
                raise RequestError(f"Filter {index + 1} with op 'in' needs a list value")
            targets = [_filter_value(tables[field["table"]], field["column"], value) for value in values]
        else:
            targets = [_filter_value(tables[field["table"]], field["column"], values)]
        prepared.append((field, operation, targets))
    output = []
    for row in fact_rows:
        passed = True
        for field, operation, targets in prepared:
            actual = _related_row(row, field, fact_table, tables, relationships, indexes)
            found = any(_matches_filter(actual, target) for target in targets)
            if (operation == "eq" and not found) or (operation == "neq" and found) or (operation == "in" and not found):
                passed = False
                break
        if passed:
            output.append(row)
    return output


def _measure_value(
    measure: dict[str, Any], rows: list[dict[str, Any]], fact_table: str,
    tables: dict[str, list[dict[str, Any]]], relationships: list[dict[str, Any]],
    indexes: dict[tuple[str, str, str], dict[str, dict[str, Any]]],
    expression_evaluator: Callable[[dict[str, Any]], Decimal | None] | None = None,
) -> tuple[Any, dict[str, int]]:
    if not isinstance(measure, dict):
        raise RequestError("Each measure must be an object")
    name = measure.get("name")
    operation = measure.get("op")
    if not isinstance(name, str) or not name.strip():
        raise RequestError("Every measure requires a non-empty 'name'")
    if not isinstance(operation, str):
        raise RequestError(f"Measure {name!r} op must be a string")
    if operation == "count_rows":
        return len(rows), {"rows_considered": len(rows), "values_used": len(rows), "blank_or_skipped": 0}
    if operation == "sum_expression":
        evaluator = expression_evaluator
        if evaluator is None:
            evaluator = _compile_expression(measure.get("expression"), set(_table_columns(tables, fact_table)))
        values = []
        skipped = 0
        for index, row in enumerate(rows, start=1):
            try:
                value = evaluator(row)
            except RequestError as exc:
                raise RequestError(f"Measure {name!r}, fact row {index}: {exc}") from exc
            if value is None:
                skipped += 1
            else:
                values.append(value)
        result = _decimal_sum(values) if values else None
        return _json_value(result), {"rows_considered": len(rows), "values_used": len(values), "blank_or_skipped": skipped}
    allowed = {"sum", "mean", "min", "max", "count_nonblank", "distinct_count"}
    if operation not in allowed:
        raise RequestError(f"Unsupported measure op {operation!r}; use {sorted(allowed | {'count_rows', 'sum_expression'})}")
    field = _field_spec(measure, f"Measure {name!r}")
    values = [_related_row(row, field, fact_table, tables, relationships, indexes) for row in rows]
    present = [value for value in values if value is not None]
    if operation == "count_nonblank":
        return len(present), {"rows_considered": len(rows), "values_used": len(present), "blank_or_skipped": len(rows) - len(present)}
    if operation == "distinct_count":
        return len({_key(value) for value in present}), {"rows_considered": len(rows), "values_used": len(present), "blank_or_skipped": len(rows) - len(present)}
    numbers = [_decimal(value, f"measure {name}") for value in present]
    if not numbers:
        result = None
    elif operation == "sum":
        result = _decimal_sum(numbers)
    elif operation == "mean":
        result = _decimal_divide(_decimal_sum(numbers), Decimal(len(numbers)))
    elif operation == "min":
        result = min(numbers)
    else:
        result = max(numbers)
    return _json_value(result), {"rows_considered": len(rows), "values_used": len(numbers), "blank_or_skipped": len(rows) - len(numbers)}


def aggregate(config: dict[str, Any], tables: dict[str, list[dict[str, Any]]]) -> dict[str, Any]:
    _validate_aggregate_operations(config)
    fact_table = config.get("fact_table")
    if not isinstance(fact_table, str) or fact_table not in tables:
        raise RequestError("aggregate requires a 'fact_table' name from the tables object")
    relationships = config.get("relationships", [])
    audit = audit_relationships(config, tables)
    if not audit["valid"]:
        return {"action": "aggregate", "valid": False, "status": "invalid", "reason": "Declared relationships failed key-integrity checks; no aggregation was run.", "model_audit": audit}
    group_by = config.get("group_by", [])
    filters = config.get("filters", [])
    measures = config.get("measures")
    if not isinstance(group_by, list) or not isinstance(filters, list) or not isinstance(measures, list) or not measures:
        raise RequestError("aggregate requires lists for group_by and filters, plus a non-empty measures list")
    for index, item in enumerate(group_by):
        field = _field_spec(item, f"group_by[{index}]")
        _validate_field_relation(field, fact_table, tables, relationships)
    measure_names = []
    compiled_expressions: dict[str, Callable[[dict[str, Any]], Decimal | None]] = {}
    fact_columns: set[str] | None = None
    for index, measure in enumerate(measures):
        if not isinstance(measure, dict):
            raise RequestError(f"measures[{index}] must be an object")
        name = measure.get("name")
        operation = measure.get("op")
        if not isinstance(name, str) or not name.strip():
            raise RequestError(f"measures[{index}] needs a non-empty 'name'")
        if name in measure_names:
            raise RequestError(f"Measure names must be unique; repeated name: {name!r}")
        measure_names.append(name)
        if operation not in _MEASURE_OPERATIONS:
            raise RequestError(f"Unsupported measures[{index}].op {operation!r}; use {sorted(_MEASURE_OPERATIONS)}")
        if operation == "sum_expression":
            if fact_columns is None:
                fact_columns = set(_table_columns(tables, fact_table))
            compiled_expressions[name] = _compile_expression(measure.get("expression"), fact_columns)
        elif operation != "count_rows":
            field = _field_spec(measure, f"Measure {name!r}")
            _validate_field_relation(field, fact_table, tables, relationships)
    indexes: dict[tuple[str, str, str], dict[str, dict[str, Any]]] = {}
    for rel in relationships:
        one_rows = tables[rel["one_table"]]
        indexes[(rel["one_table"], rel["one_column"], rel["many_table"])] = {
            _key(row[rel["one_column"]]): row for row in one_rows if row.get(rel["one_column"]) is not None
        }
    source_rows = tables[fact_table]
    filtered_rows = _apply_filters(source_rows, filters, fact_table, tables, relationships, indexes)

    def evaluate(rows: list[dict[str, Any]]):
        values, audits = {}, {}
        for measure in measures:
            result, info = _measure_value(
                measure, rows, fact_table, tables, relationships, indexes,
                expression_evaluator=compiled_expressions.get(measure["name"]),
            )
            values[measure["name"]] = result
            audits[measure["name"]] = info
        return values, audits

    totals, totals_audit = evaluate(filtered_rows)
    group_rows: list[dict[str, Any]] = []
    if group_by:
        grouped: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
        labels: dict[tuple[str, ...], list[Any]] = {}
        for row in filtered_rows:
            keys = [_related_row(row, field, fact_table, tables, relationships, indexes) for field in group_by]
            key = tuple(_key(value) for value in keys)
            grouped[key].append(row)
            labels[key] = keys
        for key in sorted(grouped, key=lambda item: tuple(str(v) for v in item)):
            values, _ = evaluate(grouped[key])
            group_rows.append({
                "group": {f"{field['table']}.{field['column']}": _json_value(value) for field, value in zip(group_by, labels[key])},
                "measures": values,
                "rows": len(grouped[key]),
            })
    return {
        "action": "aggregate",
        "valid": True,
        "status": "valid",
        "execution": "Computed by the bundled helper for the declared operations; not evaluated by Power BI or DAX.",
        "fact_table": fact_table,
        "rows_input": len(source_rows),
        "rows_used": len(filtered_rows),
        "rows_filtered_out": len(source_rows) - len(filtered_rows),
        "totals": totals,
        "measure_audit": totals_audit,
        "group_by": [{"table": field["table"], "column": field["column"]} for field in group_by],
        "groups": group_rows,
        "scope_note": "Only declared direct fact-to-dimension lookups, explicit filters, and listed aggregation operations are implemented; no DAX context, M query, or Power BI engine is present.",
    }


def run_request(config: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(config, dict):
        raise RequestError("Request root must be a JSON object", code="invalid_type", field_path="$")
    action = _validate_request_keys(config)
    _sample_policy(config)
    if action == "aggregate":
        _validate_aggregate_operations(config)
    tables = load_tables(config)
    if action == "inventory":
        return inventory(config, tables)
    if action == "model_audit":
        return audit_relationships(config, tables)
    if action == "aggregate":
        return aggregate(config, tables)
    raise RequestError("'action' must be one of: inventory, model_audit, aggregate")


def _input_files(value: Any):
    if isinstance(value, dict):
        for key, item in value.items():
            if key == "file" and isinstance(item, str):
                yield item
            else:
                yield from _input_files(item)
    elif isinstance(value, list):
        for item in value:
            yield from _input_files(item)


def ensure_output_safe(output_path: str | Path, config_path: str | Path, config: Any) -> None:
    target = Path(output_path).expanduser().resolve()
    reserved = [Path(config_path).expanduser().resolve()]
    reserved.extend(Path(filename).expanduser().resolve() for filename in _input_files(config))
    if target in set(reserved):
        raise RequestError("Output path must not overwrite the request JSON or any declared source file")
    if target.exists():
        for input_path in reserved:
            if not input_path.exists():
                continue
            try:
                aliases_input = target.samefile(input_path)
            except OSError as exc:
                raise RequestError(f"Could not safely check output file identity: {exc}") from exc
            if aliases_input:
                raise RequestError(
                    f"Output path must not overwrite the request JSON or any declared source file (same file: {input_path})"
                )


def _markdown_cell(value: Any) -> str:
    text = "(blank)" if value is None else str(value)
    return text.replace("|", "\\|").replace("\n", " ")


def _model_audit_diagnostics_markdown(result: dict[str, Any]) -> list[str]:
    lines = [
        "Status: **invalid**",
        "",
        "## Relationship diagnostics",
        "",
        "| One side | Many side | Valid | Null one-side keys | Duplicate one-side keys | Duplicate samples (opt-in) | Null many-side keys | Orphan many-side keys | Orphan samples (opt-in) |",
        "|---|---|---:|---:|---:|---|---:|---:|---|",
    ]
    for relationship in result.get("relationships", []):
        one_side = f"{relationship['one_table']}.{relationship['one_column']}"
        many_side = f"{relationship['many_table']}.{relationship['many_column']}"
        duplicate_samples = ", ".join(
            _markdown_cell(value) for value in relationship.get("duplicate_one_key_samples", [])
        ) or "—"
        orphan_samples = ", ".join(
            _markdown_cell(value) for value in relationship.get("orphan_samples", [])
        ) or "—"
        lines.append(
            f"| {_markdown_cell(one_side)} | {_markdown_cell(many_side)} | {relationship['valid']} "
            f"| {relationship['one_key_nulls']} | {relationship['duplicate_one_key_count']} "
            f"| {duplicate_samples} | {relationship['many_key_nulls']} | {relationship['orphan_count']} | {orphan_samples} |"
        )
    for warning in result.get("warnings", []):
        lines.extend(["", f"Warning: {_markdown_cell(warning)}"])
    if result.get("scope_note"):
        lines.extend(["", str(result["scope_note"])])
    return lines


def format_markdown(result: dict[str, Any]) -> str:
    action = result.get("action", "analysis")
    lines = [f"# Power BI-style {action.replace('_', ' ')}", "", "> Helper output; not Power BI rendering or DAX evaluation.", ""]
    if not result.get("valid", False):
        invalid_reason = result.get("reason") or result.get("error")
        if not invalid_reason and action == "model_audit" and isinstance(result.get("relationships"), list):
            invalid_reason = "Declared relationship integrity checks failed; see the diagnostics below."
        lines.extend(["## RESULT INVALID", "", str(invalid_reason or "Validation failed."), ""])
        audit = result.get("model_audit")
        if audit:
            lines.append(_json_dumps(audit, indent=2))
        elif action == "model_audit" and isinstance(result.get("relationships"), list):
            lines.extend(_model_audit_diagnostics_markdown(result))
        return "\n".join(lines).rstrip() + "\n"
    if action == "inventory":
        for table, info in result["tables"].items():
            lines += [f"## {table}", "", f"Rows: {info['rows']} · Columns: {info['columns']} · Duplicate rows: {info['duplicate_rows']}", "", "| Column | Inferred type | Nonmissing | Missing | Distinct | Sample values |", "|---|---|---:|---:|---:|---|"]
            for field in info["fields"]:
                samples = ", ".join(_markdown_cell(value) for value in field["sample_values"])
                lines.append(f"| {_markdown_cell(field['name'])} | {_markdown_cell(field['inferred_type'])} | {field['nonmissing']} | {field['missing']} | {field['distinct']} | {samples} |")
            lines.append("")
    elif action == "model_audit":
        lines += [f"Status: **{result['status']}**", "", "| One side | Many side | Valid | Duplicate one keys | Orphan many keys | Null many keys |", "|---|---|---:|---:|---:|---:|"]
        for rel in result["relationships"]:
            lines.append(f"| {rel['one_table']}.{rel['one_column']} | {rel['many_table']}.{rel['many_column']} | {rel['valid']} | {rel['duplicate_one_key_count']} | {rel['orphan_count']} | {rel['many_key_nulls']} |")
        for warning in result.get("warnings", []):
            lines += ["", f"Warning: {warning}"]
        lines += ["", result["scope_note"]]
    elif action == "aggregate":
        lines += [f"Fact table: `{result['fact_table']}` · Rows used: {result['rows_used']} / {result['rows_input']} · Filtered out: {result['rows_filtered_out']}", "", "## Totals", "", "| Measure | Value | Values used | Blank/skipped |", "|---|---:|---:|---:|"]
        for measure, value in result["totals"].items():
            audit = result["measure_audit"][measure]
            lines.append(f"| {_markdown_cell(measure)} | {_markdown_cell(value)} | {audit['values_used']} | {audit['blank_or_skipped']} |")
        if result["groups"]:
            measure_names = list(result["totals"])
            group_names = [f"{field['table']}.{field['column']}" for field in result["group_by"]]
            lines += ["", "## Groups", "", "| " + " | ".join(group_names + measure_names + ["Rows"]) + " |", "|" + "---|" * (len(group_names) + len(measure_names) + 1)]
            for group in result["groups"]:
                cells = [group["group"].get(name.rsplit(".", 1)[0] + "." + name.rsplit(".", 1)[-1]) for name in group_names]
                cells += [group["measures"].get(name) for name in measure_names] + [group["rows"]]
                lines.append("| " + " | ".join(_markdown_cell(value) for value in cells) + " |")
        lines += ["", result["scope_note"]]
    return "\n".join(lines).rstrip() + "\n"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", required=True, help="UTF-8 JSON request file")
    parser.add_argument("--output", help="Output path; defaults to stdout")
    parser.add_argument("--format", choices=("json", "markdown"), default="json")
    args = parser.parse_args(argv)
    config_path = Path(args.config).expanduser()
    target = Path(args.output).expanduser() if args.output else None
    if target and target.resolve() == config_path.resolve():
        print("Error: output path must not overwrite the request JSON", file=sys.stderr)
        return 2
    try:
        with config_path.open("r", encoding="utf-8") as handle:
            config = load_json(handle)
        if target:
            ensure_output_safe(target, config_path, config)
        result = run_request(config)
    except (OSError, json.JSONDecodeError, RequestError) as exc:
        result = {"valid": False, "status": "error", "error": str(exc)}
        if isinstance(exc, RequestError):
            result["error_code"] = exc.code
            if exc.field_path is not None:
                result["error_path"] = exc.field_path
        elif isinstance(exc, json.JSONDecodeError):
            result["error_code"] = "invalid_json"
        else:
            result["error_code"] = "input_error"
        print(f"Error: {exc}", file=sys.stderr)
        exit_code = 2
    else:
        exit_code = 0
    output = format_markdown(result) if args.format == "markdown" else _json_dumps(result, indent=2)
    if target and exit_code == 0:
        try:
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_text(output + ("" if output.endswith("\n") else "\n"), encoding="utf-8")
        except (OSError, UnicodeError) as exc:
            print(f"Error writing output {target}: {exc}", file=sys.stderr)
            return 2
    else:
        print(output)
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main())
